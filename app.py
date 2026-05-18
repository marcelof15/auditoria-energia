from flask import Flask, request, jsonify, send_from_directory, send_file
from flask_cors import CORS
import os, tempfile, traceback, io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from parser import processar_pdf

app = Flask(__name__, static_folder='static')
CORS(app)

@app.route('/')
def index():
    return send_from_directory('static', 'index.html')

@app.route('/processar', methods=['POST'])
def processar():
    try:
        files = request.files.getlist('faturas')
        if not files:
            return jsonify({'erro': 'Nenhum arquivo enviado'}), 400
        resultados = []
        for f in files:
            if not f.filename.lower().endswith('.pdf'):
                continue
            tmp = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False)
            f.save(tmp.name); tmp.close()
            try:
                dados = processar_pdf(tmp.name)
                dados['arquivo'] = f.filename
                resultados.append(dados)
            except Exception as e:
                resultados.append({'arquivo': f.filename, 'tipo': 'erro', 'erro': str(e)})
            finally:
                os.unlink(tmp.name)
        pares = cruzar_pares(resultados)
        return jsonify({'faturas': resultados, 'pares': pares})
    except Exception as e:
        traceback.print_exc()
        return jsonify({'erro': str(e)}), 500

@app.route('/exportar', methods=['POST'])
def exportar():
    try:
        data   = request.get_json()
        pares  = data.get('pares', [])
        desc_esp = data.get('desc_esp', 20)
        wb  = gerar_excel(pares, desc_esp)
        buf = io.BytesIO()
        wb.save(buf); buf.seek(0)
        return send_file(buf,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True, download_name='auditoria_energia.xlsx')
    except Exception as e:
        traceback.print_exc()
        return jsonify({'erro': str(e)}), 500

# ── HELPERS ──────────────────────────────────────────────────────────
def _fill(hex): return PatternFill('solid', fgColor=hex)
def _font(bold=False, color='000000', size=10, italic=False):
    return Font(name='Arial', bold=bold, color=color, size=size, italic=italic)
def _border():
    s = Side(style='thin', color='BBBBBB')
    return Border(top=s, bottom=s, left=s, right=s)
def _center(): return Alignment(horizontal='center', vertical='center', wrap_text=True)
def _right():  return Alignment(horizontal='right',  vertical='center')
def _left(w=False): return Alignment(horizontal='left', vertical='center', wrap_text=w)
BRL = '#,##0.00'; PCT = '0.0%'; INT = '#,##0'

def _set(ws, row, col, value, bold=False, color='000000', size=10, bg=None,
         fmt=None, align=None, italic=False):
    c = ws.cell(row=row, column=col, value=value)
    c.font   = _font(bold, color, size, italic)
    c.border = _border()
    if bg:    c.fill  = _fill(bg)
    if fmt:   c.number_format = fmt
    if align: c.alignment = align
    return c

# ── GERADOR EXCEL ────────────────────────────────────────────────────
def gerar_excel(pares, desc_esp=20):
    wb = Workbook()
    pv = [p for p in pares if not p.get('sem_cosern') and not p.get('sem_origo')]

    total_ref    = sum(p.get('ref_origo',0)    for p in pv)
    total_cosern = sum(p.get('total_cosern',0) for p in pv)
    total_origo  = sum(p.get('total_origo',0)  for p in pv)
    total_pago   = sum(p.get('total_pago',0)   for p in pv)
    total_eco    = sum(p.get('economia',0)     for p in pv)
    desc_medio   = (total_eco / total_ref * 100) if total_ref else 0
    comps = sorted({p.get('competencia','') for p in pv if p.get('competencia')})

    # ═══════════════════════════════════════════════════
    # ABA 1 — RESUMO
    # ═══════════════════════════════════════════════════
    ws = wb.active; ws.title = 'Resumo'
    ws.sheet_properties.tabColor = '1a3560'

    # Título
    ws.merge_cells('A1:L1')
    c = ws.cell(1,1,'⚡ MB ENERGIA — AUDITORIA DE FATURAS COSERN + ORIGO')
    c.font = _font(True,'FFFFFF',13); c.fill = _fill('1a3560'); c.alignment = _center()
    ws.row_dimensions[1].height = 32

    ws.merge_cells('A2:L2')
    c = ws.cell(2,1,f"Competência(s): {', '.join(comps)}   |   Desconto contratado: {desc_esp:.0f}%   |   {len(pv)} par(es) analisado(s)")
    c.font = _font(italic=True,size=9,color='555555'); c.alignment = _left()
    ws.row_dimensions[2].height = 14; ws.row_dimensions[3].height = 8

    # Cards métricas
    g_ok = desc_medio >= desc_esp * 0.95
    cards = [
        ('A','D','Ref. bruta COSERN (sem GD)',total_ref,BRL,'1B4F72'),
        ('E','G','Total pago (COSERN+Origo)', total_pago,BRL,'1a3560'),
        ('H','J','Economia total',             total_eco, BRL,'1E8449' if total_eco>=0 else 'C0392B'),
        ('K','L','Desconto médio real',         desc_medio/100,PCT,'1E8449' if g_ok else 'C0392B'),
    ]
    for s,e,lbl,val,fmt,color in cards:
        sc=ord(s)-64; ec=ord(e)-64
        ws.merge_cells(f'{s}4:{e}4'); ws.merge_cells(f'{s}5:{e}5')
        for col in range(sc,ec+1):
            for row in [4,5]:
                ws.cell(row,col).fill=_fill('F0F4FF'); ws.cell(row,col).border=_border()
        c4=ws.cell(4,sc,lbl); c4.font=_font(size=9,color='555555'); c4.alignment=_center()
        c5=ws.cell(5,sc,val); c5.font=_font(True,color,15); c5.number_format=fmt; c5.alignment=_center()
    ws.row_dimensions[4].height=16; ws.row_dimensions[5].height=28; ws.row_dimensions[6].height=8

    # Cabeçalho tabela
    hdrs=['#','Instalação','UC','Competência','Base desconto (R$)','Pago COSERN (R$)',
          'Pago Origo (R$)','Total pago (R$)','Economia (R$)','Desc. real (%)','Esp. (%)','Status']
    wids=[4,28,10,12,16,16,16,16,14,12,10,16]
    for i,(h,w) in enumerate(zip(hdrs,wids)):
        ws.column_dimensions[get_column_letter(i+1)].width=w
        c=ws.cell(7,i+1,h); c.font=_font(True,'FFFFFF',9)
        c.fill=_fill('1a3560'); c.alignment=_center(); c.border=_border()
    ws.row_dimensions[7].height=32

    for ri,p in enumerate(pv,1):
        row=7+ri; ws.row_dimensions[row].height=18
        st=p.get('status','')
        bg='E8F5E9' if st=='conforme' else 'FFF8E1' if st=='atencao' else 'FFEBEE'
        st_txt='✅ Conforme' if st=='conforme' else '⚠ Atenção' if st=='atencao' else '✕ Não conforme'
        st_c='1E8449' if st=='conforme' else '856404' if st=='atencao' else 'C0392B'
        eco_c='1E8449' if p.get('economia',0)>=0 else 'C0392B'
        dr_c='1E8449' if p.get('desc_real',0)>=desc_esp*0.95 else 'C0392B'
        rows=[
            (ri,None,False,'000000'),(p.get('instalacao',''),None,True,'000000'),
            (p.get('uc',''),None,False,'000000'),(p.get('competencia',''),None,False,'000000'),
            (p.get('ref_origo',0),BRL,False,'000000'),(p.get('total_cosern',0),BRL,False,'000000'),
            (p.get('total_origo',0),BRL,False,'000000'),(p.get('total_pago',0),BRL,True,'000000'),
            (p.get('economia',0),BRL,True,eco_c),(p.get('desc_real',0)/100,PCT,True,dr_c),
            (desc_esp/100,PCT,False,'555555'),(st_txt,None,True,st_c),
        ]
        aligns=[_center(),_left(),_center(),_center()]+[_right()]*8
        for ci,((v,fmt,bold,color),align) in enumerate(zip(rows,aligns),1):
            c=ws.cell(row,ci,v); c.fill=_fill(bg); c.border=_border()
            c.font=_font(bold,color,10); c.alignment=align
            if fmt: c.number_format=fmt

    # Totais
    tr=7+len(pv)+1; ws.row_dimensions[tr].height=20
    ws.merge_cells(f'A{tr}:D{tr}')
    c=ws.cell(tr,1,'TOTAIS'); c.font=_font(True,'FFFFFF',10)
    c.fill=_fill('1a3560'); c.alignment=_center(); c.border=_border()
    for ci,v,fmt in [(5,total_ref,BRL),(6,total_cosern,BRL),(7,total_origo,BRL),
                     (8,total_pago,BRL),(9,total_eco,BRL),(10,desc_medio/100,PCT)]:
        c=ws.cell(tr,ci,v); c.font=_font(True,'FFFFFF',10)
        c.fill=_fill('1a3560'); c.number_format=fmt; c.alignment=_right(); c.border=_border()
    for ci in [11,12]:
        ws.cell(tr,ci).fill=_fill('1a3560'); ws.cell(tr,ci).border=_border()
    ws.freeze_panes='A8'

    # ═══════════════════════════════════════════════════
    # ABA 2 — DETALHAMENTO
    # ═══════════════════════════════════════════════════
    ws2 = wb.create_sheet('Detalhamento')
    ws2.sheet_properties.tabColor = '1E8449'
    ws2.merge_cells('A1:AB1')
    c=ws2.cell(1,1,'📋 DETALHAMENTO COMPLETO POR FATURA')
    c.font=_font(True,'FFFFFF',12); c.fill=_fill('1E3A2E'); c.alignment=_center()
    ws2.row_dimensions[1].height=28

    d_hdrs=['#','Instalação','UC','Torre','Competência','Consumo kWh',
            'TUSD','TE','Bandeira','IP','Parcelas',
            'G1 kWh','G1 (R$)','G2 kWh','G2 (R$)',
            'Base desc.','Pago COSERN','Base ICMS',
            'Origo bruto','PIS/COFINS','Desc. Comerciais','Cobranças Adicionais',
            'Total Origo','Total pago','Economia','Desc. real %','Esp. %','Status']
    d_wids=[4,24,10,10,12,11,12,12,11,9,11,10,12,10,12,13,14,12,13,12,16,18,13,13,12,11,9,15]
    d_fmts=[None,None,None,None,None,INT,BRL,BRL,BRL,BRL,BRL,INT,BRL,INT,BRL,BRL,BRL,BRL,BRL,BRL,BRL,BRL,BRL,BRL,BRL,PCT,PCT,None]

    for i,(h,w) in enumerate(zip(d_hdrs,d_wids)):
        ws2.column_dimensions[get_column_letter(i+1)].width=w
        c=ws2.cell(2,i+1,h); c.font=_font(True,'FFFFFF',9)
        c.fill=_fill('1E3A2E'); c.alignment=_center(); c.border=_border()
    ws2.row_dimensions[2].height=36

    for ri,p in enumerate(pv,1):
        row=2+ri; ws2.row_dimensions[row].height=18
        st=p.get('status','')
        bg='F1FFF5' if st=='conforme' else 'FFFDE7' if st=='atencao' else 'FFF0F0'
        st_txt='✅ Conforme' if st=='conforme' else '⚠ Atenção' if st=='atencao' else '✕ Não conforme'
        vals=[ri,p.get('instalacao',''),p.get('uc',''),p.get('torre',''),p.get('competencia',''),
              p.get('consumo_kwh',0),p.get('tusd',0),p.get('te',0),p.get('bandeira',0),
              p.get('ip',0),p.get('parcelas',0),p.get('g1_kwh',0),p.get('g1_credito',0),
              p.get('g2_kwh',0),p.get('g2_credito',0),p.get('ref_origo',0),
              p.get('total_cosern',0),p.get('base_icms',0),p.get('origo_bruto',0),
              p.get('pis_cofins',0),p.get('desc_comerciais',0),p.get('cobrancas_adic',0),
              p.get('total_origo',0),p.get('total_pago',0),p.get('economia',0),
              p.get('desc_real',0)/100,desc_esp/100,st_txt]
        for ci,(v,fmt) in enumerate(zip(vals,d_fmts),1):
            c=ws2.cell(row,ci,v); c.fill=_fill(bg); c.border=_border()
            c.font=_font(size=9); c.alignment=_right() if ci>5 else (_center() if ci==1 else _left())
            if fmt: c.number_format=fmt
    ws2.freeze_panes='A3'

    # ═══════════════════════════════════════════════════
    # ABA 3 — ALERTAS
    # ═══════════════════════════════════════════════════
    ws3=wb.create_sheet('Alertas'); ws3.sheet_properties.tabColor='C0392B'
    ws3.merge_cells('A1:G1')
    c=ws3.cell(1,1,'⚠ ALERTAS — Instalações com desconto abaixo do esperado')
    c.font=_font(True,'FFFFFF',11); c.fill=_fill('C0392B'); c.alignment=_center()
    ws3.row_dimensions[1].height=28

    a_hdrs=['Instalação','UC','Competência','Desc. real (%)','Esp. (%)','Diferença (pp)','Situação']
    a_wids=[30,10,12,14,12,14,20]
    for i,(h,w) in enumerate(zip(a_hdrs,a_wids)):
        ws3.column_dimensions[get_column_letter(i+1)].width=w
        c=ws3.cell(2,i+1,h); c.font=_font(True,'FFFFFF',10)
        c.fill=_fill('C0392B'); c.alignment=_center(); c.border=_border()
    ws3.row_dimensions[2].height=22

    alertas=[p for p in pv if p.get('status')!='conforme']
    if not alertas:
        ws3.merge_cells('A3:G3')
        c=ws3.cell(3,1,'✅ Nenhum alerta — todas as instalações estão conformes!')
        c.font=_font(True,'1E8449',11); c.fill=_fill('E8F5E9'); c.alignment=_center()
        ws3.row_dimensions[3].height=24
    else:
        for ri,p in enumerate(alertas,1):
            row=2+ri; ws3.row_dimensions[row].height=18
            st=p.get('status','')
            bg='FFF8E1' if st=='atencao' else 'FFEBEE'
            st_txt='⚠ Atenção' if st=='atencao' else '✕ Não conforme'
            st_c='856404' if st=='atencao' else 'C0392B'
            dr=p.get('desc_real',0); dif=dr-desc_esp
            vals=[p.get('instalacao',''),p.get('uc',''),p.get('competencia',''),
                  dr/100,desc_esp/100,dif/100,st_txt]
            fmts=[None,None,None,PCT,PCT,'+0.0%;-0.0%;0.0%',None]
            for ci,(v,fmt) in enumerate(zip(vals,fmts),1):
                c=ws3.cell(row,ci,v); c.fill=_fill(bg); c.border=_border()
                c.font=_font(bold=(ci==7),color=(st_c if ci==7 else '000000'),size=10)
                c.alignment=_right() if ci in [4,5,6] else _left()
                if fmt: c.number_format=fmt
    return wb

def cruzar_pares(faturas):
    cosern=[f for f in faturas if f.get('tipo')=='cosern']
    origo=[f for f in faturas if f.get('tipo')=='origo']
    pares=[]; used_c=set(); used_o=set()
    for oi,o in enumerate(origo):
        for ci,c in enumerate(cosern):
            if ci in used_c: continue
            if (c.get('uc') and o.get('uc') and c['uc']==o['uc']) or c.get('competencia')==o.get('competencia'):
                pares.append(calcular_auditoria(c,o)); used_c.add(ci); used_o.add(oi); break
        else:
            if oi not in used_o:
                pares.append({'sem_cosern':True,'origo':o,'competencia':o.get('competencia',''),
                              'instalacao':o.get('cliente','')+'  UC '+o.get('uc','')})
    for ci,c in enumerate(cosern):
        if ci not in used_c:
            pares.append({'sem_origo':True,'cosern':c,'competencia':c.get('competencia',''),
                          'instalacao':c.get('cliente','')+'  UC '+c.get('uc','')})
    return pares

def calcular_auditoria(c,o,desc_esp=0.20):
    ref_total=(c.get('tusd',0)+c.get('te',0)+c.get('bandeira',0)+
               c.get('iluminacao_publica',0)+c.get('parcelas',0))
    g2=c.get('g2_credito',0); ref_origo=ref_total-g2
    tc=c.get('total_fatura',0); to=o.get('total_origo',0); tp=tc+to
    eco=ref_origo-tp; dr=(eco/ref_origo*100) if ref_origo else 0
    diff=dr-(desc_esp*100)
    st='conforme' if diff>=-1 else 'atencao' if diff>=-5 else 'nao_conforme'
    torre=c.get('torre',''); cliente=c.get('cliente','') or o.get('cliente','')
    return {'instalacao':f"{cliente} — Torre {torre}" if torre else cliente,
            'uc':c.get('uc','') or o.get('uc',''),'torre':torre,
            'competencia':c.get('competencia','') or o.get('competencia',''),
            'consumo_kwh':c.get('consumo_kwh',0),'ref_total':round(ref_total,2),
            'ref_origo':round(ref_origo,2),'tusd':c.get('tusd',0),'te':c.get('te',0),
            'bandeira':c.get('bandeira',0),'ip':c.get('iluminacao_publica',0),
            'parcelas':c.get('parcelas',0),'g1_credito':round(c.get('g1_credito',0),2),
            'g1_kwh':c.get('g1_kwh',0),'g2_credito':round(g2,2),'g2_kwh':c.get('g2_kwh',0),
            'total_cosern':round(tc,2),'origo_bruto':o.get('origo_bruto',0),
            'pis_cofins':o.get('pis_cofins',0),'desc_comerciais':o.get('descontos_comerciais',0),
            'cobrancas_adic':o.get('cobrancas_adicionais',0),'total_origo':round(to,2),
            'total_pago':round(tp,2),'economia':round(eco,2),'desc_real':round(dr,2),
            'desc_esp':desc_esp*100,'base_icms':c.get('base_icms',0),'status':st,
            'arquivo_cosern':c.get('arquivo',''),'arquivo_origo':o.get('arquivo','')}

if __name__ == '__main__':
    os.makedirs('static', exist_ok=True)
    app.run(debug=True, port=5000)
